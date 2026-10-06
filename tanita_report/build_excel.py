"""產生 Excel 驗算檔：主要統計量以 Excel 內建公式重算，並列出 Python 計算值供比對。

由 build_report.py 呼叫：build(D, res, out_path)
另產生同名 .sps（SPSS 語法），可在 SPSS 讀入同一份 Excel 重跑檢定。
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

import analyze as A

FONT = "Microsoft JhengHei"
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")          # 原始資料（硬寫數值）
F_PY = Font(name=FONT, size=10, color="7F7F7F", italic=True)  # Python 參考值
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_SUB = Font(name=FONT, size=9, color="595959")
F_TITLE = Font(name=FONT, size=14, bold=True, color="04706F")
F_BOLD = Font(name=FONT, size=10, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="069494")
FILL_SUB = PatternFill("solid", fgColor="E4F4F3")
FILL_PY = PatternFill("solid", fgColor="F2F2F2")
FILL_CHK = PatternFill("solid", fgColor="FFFBE2")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

BIN_FORMULA = {  # 以同列欄位計算分類旗標的公式樣板（{bmi}{fat}{vf}{sex} 會替換成儲存格）
    "ow": "=IF({bmi}>=24,1,0)",
    "ob": "=IF({bmi}>=27,1,0)",
    "fat": '=IF(OR(AND({sex}="M",{fat}>=25),AND({sex}="F",{fat}>=30)),1,0)',
    "vfat": "=IF({vf}>=10,1,0)",
}
BMI_CAT_F = '=IF({bmi}<18.5,"過輕",IF({bmi}<24,"正常",IF({bmi}<27,"過重","肥胖")))'
VFAT_CAT_F = '=IF({vf}<10,"{c0}",IF({vf}<15,"{c1}","{c2}"))'


def head(ws, row, labels, widths=None):
    for j, lab in enumerate(labels, 1):
        c = ws.cell(row, j, lab)
        c.font, c.fill, c.alignment = F_HEAD, FILL_HEAD, CENTER
    ws.row_dimensions[row].height = 32
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[L(j)].width = w


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB


def style_range(ws, r0, r1, c0, c1, font=F_BASE, fmt=None, fill=None):
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            cell = ws.cell(r, c)
            cell.font = font
            cell.border = THIN
            if fmt:
                cell.number_format = fmt
            if fill:
                cell.fill = fill


def build(D, res, out_path):
    pre, post, ids = D["pre"], D["post"], D["ids"]
    fields = {**A.MAIN, **A.SEG}
    wb = Workbook()

    # ------------------------------------------------------------ 說明
    ws = wb.active
    ws.title = "說明"
    info = res["info"]
    rows = [
        ("TANITA 身體組成前後測 — Excel 驗算檔", None),
        (f"職場：{info['workplace']}　前測 {info['pre_date']}　後測 {info['post_date']}　完成前後測 {info['n_paired']} 人", None),
        ("", None),
        ("工作表", "內容"),
        ("欄位對照", "TANITA MC-780 原始欄位代碼、中文名稱、單位"),
        ("配對資料", "完成前後測者的去識別化資料（序號取代編號），差值與分類旗標皆為公式"),
        ("前測資料", "前測全部受測者（卡方檢定、流失分析使用），分類旗標為公式"),
        ("配對檢定", "全體：平均、標準差、配對 t、95% 信賴區間、效果量 — 全為 Excel 公式；並列 Python 值與差異"),
        ("分層檢定", "性別、年齡層各分層的配對 t 檢定（SUMIF / SUMPRODUCT 公式）"),
        ("組間比較", "性別 Welch t、年齡層單因子 ANOVA（公式）；Mann-Whitney、Kruskal-Wallis 為 Python 值"),
        ("McNemar", "風險分類前後變化：McNemar 精確檢定與 Bowker 對稱性檢定（公式）"),
        ("卡方檢定", "前測風險分類 × 性別／年齡層：觀察次數、期望次數、χ²、p 值（公式）；Fisher 精確檢定為 Python 值"),
        ("流失分析", "完成前後測者 vs 僅前測者的前測值比較（Welch t 公式、卡方公式）"),
        ("", None),
        ("顏色說明", ""),
        ("藍色字", "原始資料（直接輸入的數值）"),
        ("黑色字", "Excel 公式計算結果"),
        ("灰色斜體（灰底）", "Python（SciPy）計算的參考值，用於交叉驗證"),
        ("黃底", "Excel − Python 差異；應接近 0（< 0.0001）"),
        ("", None),
        ("無法以 Excel 公式計算的項目", "Wilcoxon 符號等級檢定、Shapiro-Wilk 常態檢定、Mann-Whitney U、Kruskal-Wallis、Fisher 精確檢定、Holm 校正 — 列為 Python 參考值；可用同名 .sps 語法檔在 SPSS 重跑驗證。"),
        ("單位換算", "基礎代謝量原始資料為 kJ，已 ÷ 4.184 換算為 kcal 並取整數。"),
        ("資料處理", "；".join(info["cleaning_log"]) or "無"),
        ("隱私", "本檔含個人層級（去識別化）資料，僅供分析人員驗算使用，請勿提供給職場或公開。"),
    ]
    for i, (a, b) in enumerate(rows, 1):
        ws.cell(i, 1, a).font = F_BASE
        if b is not None:
            ws.cell(i, 2, b).font = F_BASE
            ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = F_TITLE
    ws["A2"].font = F_SUB
    for r in (4, 15):
        ws.cell(r, 1).font = F_BOLD
        ws.cell(r, 2).font = F_BOLD
    ws["A16"].font = F_INPUT
    ws["A18"].font, ws["A18"].fill = F_PY, FILL_PY
    ws["A19"].fill = FILL_CHK
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100

    # ------------------------------------------------------------ 欄位對照
    ws = wb.create_sheet("欄位對照")
    head(ws, 1, ["原始欄位", "中文名稱", "單位", "說明", "納入統計"], [14, 26, 10, 60, 10])
    r = 2
    for code, (name, unit, note) in A.OTHER_FIELDS.items():
        for j, v in enumerate([code, name, unit, note, "否"], 1):
            ws.cell(r, j, v).font = F_BASE
        r += 1
    for code, meta in fields.items():
        for j, v in enumerate([code, meta[0], meta[1], meta[4] if len(meta) > 4 else "", "是"], 1):
            ws.cell(r, j, v).font = F_BASE
        r += 1

    # ------------------------------------------------------------ 配對資料
    ws = wb.create_sheet("配對資料")
    codes = list(fields)
    hdr1 = ["序號", "性別", "前測年齡", "年齡層"]
    hdr2 = ["ID", "SEX", "AGE", "AGEGRP"]
    col = {}  # code -> (pre col, post col, diff col)
    c = 5
    for code in codes:
        name = fields[code][0]
        col[code] = (c, c + 1, c + 2)
        hdr1 += [f"{name}\n前測", f"{name}\n後測", f"{name}\n差值"]
        hdr2 += [f"{code}_pre", f"{code}_post", f"{code}_d"]
        c += 3
    flag_cols = {}
    for bc in A.BINARY:
        flag_cols[bc] = (c, c + 1)
        hdr1 += [f"{A.BINARY[bc][0]}\n前測", f"{A.BINARY[bc][0]}\n後測"]
        hdr2 += [f"{bc}_pre", f"{bc}_post"]
        c += 2
    cat_cols = {"bmi": (c, c + 1), "vfat": (c + 2, c + 3)}
    hdr1 += ["BMI 分級\n前測", "BMI 分級\n後測", "內臟脂肪分級\n前測", "內臟脂肪分級\n後測"]
    hdr2 += ["BMIcat_pre", "BMIcat_post", "VFcat_pre", "VFcat_post"]
    head(ws, 1, hdr1)
    for j, v in enumerate(hdr2, 1):
        ws.cell(2, j, v).font = F_SUB
        ws.cell(2, j).fill = FILL_SUB
    ws.row_dimensions[1].height = 44
    R0 = 3
    n = len(ids)
    R1 = R0 + n - 1
    for i, pid in enumerate(ids):
        rr = R0 + i
        a, b = pre.loc[pid], post.loc[pid]
        for j, v in enumerate([i + 1, a.SEX, int(a.AGE), a.AGEGRP], 1):
            ws.cell(rr, j, v).font = F_INPUT
        for code, (cp, cq, cd) in col.items():
            ws.cell(rr, cp, float(a[code])).font = F_INPUT
            ws.cell(rr, cq, float(b[code])).font = F_INPUT
            ws.cell(rr, cd, f"={L(cq)}{rr}-{L(cp)}{rr}").font = F_BASE
        sexc = f"$B{rr}"
        for bc, (cpre, cpost) in flag_cols.items():
            for cc, idx in ((cpre, 0), (cpost, 1)):
                ref = {"bmi": f"{L(col['BMI'][idx])}{rr}", "fat": f"{L(col['FATP'][idx])}{rr}",
                       "vf": f"{L(col['VFATL'][idx])}{rr}", "sex": sexc}
                ws.cell(rr, cc, BIN_FORMULA[bc].format(**ref)).font = F_BASE
        for idx in (0, 1):
            ws.cell(rr, cat_cols["bmi"][idx], BMI_CAT_F.format(bmi=f"{L(col['BMI'][idx])}{rr}")).font = F_BASE
            ws.cell(rr, cat_cols["vfat"][idx], VFAT_CAT_F.format(vf=f"{L(col['VFATL'][idx])}{rr}",
                                                                 c0=A.VFAT_CATS[0], c1=A.VFAT_CATS[1], c2=A.VFAT_CATS[2])).font = F_BASE
    for j in range(1, len(hdr1) + 1):
        ws.column_dimensions[L(j)].width = 11
    ws.freeze_panes = "E3"
    PD = "配對資料"
    rng = lambda cc: f"{PD}!${L(cc)}${R0}:${L(cc)}${R1}"
    SEXR, AGER = f"{PD}!$B${R0}:$B${R1}", f"{PD}!$D${R0}:$D${R1}"

    # ------------------------------------------------------------ 前測資料（全部前測者）
    wp = wb.create_sheet("前測資料")
    base_codes = ["WEIGHT", "BMI", "FATP", "PMM", "VFATL"]
    h1 = ["序號", "性別", "年齡", "年齡層", "完成後測(1=是)"] + [fields[c][0] for c in base_codes] + [A.BINARY[b][0] for b in A.BINARY]
    h2 = ["ID", "SEX", "AGE", "AGEGRP", "COMPLETED"] + base_codes + list(A.BINARY)
    head(wp, 1, h1)
    for j, v in enumerate(h2, 1):
        wp.cell(2, j, v).font = F_SUB
        wp.cell(2, j).fill = FILL_SUB
    wp.row_dimensions[1].height = 44
    bcol = {c: 6 + i for i, c in enumerate(base_codes)}
    fcol = {b: 6 + len(base_codes) + i for i, b in enumerate(A.BINARY)}
    B0 = 3
    for i, (pid, a) in enumerate(pre.sort_values("DATETIME").iterrows()):
        rr = B0 + i
        for j, v in enumerate([i + 1, a.SEX, int(a.AGE), a.AGEGRP, int(pid in ids)], 1):
            wp.cell(rr, j, v).font = F_INPUT
        for cname, cc in bcol.items():
            wp.cell(rr, cc, float(a[cname])).font = F_INPUT
        ref = {"bmi": f"{L(bcol['BMI'])}{rr}", "fat": f"{L(bcol['FATP'])}{rr}", "vf": f"{L(bcol['VFATL'])}{rr}", "sex": f"$B{rr}"}
        for bc, cc in fcol.items():
            wp.cell(rr, cc, BIN_FORMULA[bc].format(**ref)).font = F_BASE
    B1 = B0 + len(pre) - 1
    for j in range(1, len(h1) + 1):
        wp.column_dimensions[L(j)].width = 12
    wp.freeze_panes = "B3"
    BD = "前測資料"
    brng = lambda cc: f"{BD}!${L(cc)}${B0}:${L(cc)}${B1}"

    # ------------------------------------------------------------ 配對檢定
    ws = wb.create_sheet("配對檢定")
    title(ws, "配對 t 檢定（全體）", "C–S 欄為 Excel 公式；T 欄起為 Python 參考值與比對。判定：差值 Shapiro-Wilk p ≥ 0.05 用配對 t，否則用 Wilcoxon。")
    H = ["指標", "單位", "n", "前測平均", "前測SD", "後測平均", "後測SD", "平均差\n(後−前)", "差值SD", "標準誤SE",
         "t", "df", "p 值\n(TDIST)", "t 臨界值\n(0.05)", "95% CI\n下限", "95% CI\n上限", "p 值驗證\n(TTEST)", "效果量\nCohen's dz", "變化率",
         "Python\n配對 t p", "差異\nExcel−Python", "Python\nWilcoxon p", "Python\nShapiro-Wilk p", "判定依據", "判定 p", "Python\nHolm 校正 p"]
    head(ws, 4, H, [24, 8, 6, 10, 9, 10, 9, 10, 9, 9, 8, 5, 10, 9, 10, 10, 10, 10, 9, 10, 11, 10, 11, 9, 9, 10])
    allm = res["strata"]["all"]["main"]
    allseg = res["strata"]["all"]["seg"]
    r = 5
    for code in codes:
        s = allm.get(code) or allseg.get(code)
        cp, cq, cd = col[code]
        P, Q, Dd = rng(cp), rng(cq), rng(cd)
        vals = [fields[code][0], fields[code][1], f"=COUNT({Dd})", f"=AVERAGE({P})", f"=STDEV({P})", f"=AVERAGE({Q})", f"=STDEV({Q})",
                f"=AVERAGE({Dd})", f"=STDEV({Dd})", f"=I{r}/SQRT(C{r})", f"=IFERROR(H{r}/J{r},0)", f"=C{r}-1",
                f"=IFERROR(TDIST(ABS(K{r}),L{r},2),1)", f"=TINV(0.05,L{r})", f"=H{r}-N{r}*J{r}", f"=H{r}+N{r}*J{r}",
                f"=IFERROR(TTEST({P},{Q},2,1),1)", f"=IFERROR(H{r}/I{r},0)", f"=H{r}/D{r}",
                s.get("p_t"), f"=M{r}-T{r}", s.get("p_w"), s.get("p_sw"), s.get("test"),
                f'=IF(X{r}="配對t",M{r},V{r})', s.get("p_holm")]
        for j, v in enumerate(vals, 1):
            ws.cell(r, j, v)
        r += 1
    R_end = r - 1
    style_range(ws, 5, R_end, 1, 26)
    style_range(ws, 5, R_end, 4, 11, fmt="0.0000")
    style_range(ws, 5, R_end, 13, 18, fmt="0.0000")
    style_range(ws, 5, R_end, 19, 19, fmt="0.00%")
    style_range(ws, 5, R_end, 20, 20, font=F_PY, fmt="0.000000", fill=FILL_PY)
    style_range(ws, 5, R_end, 21, 21, fmt="0.0E+00", fill=FILL_CHK)
    style_range(ws, 5, R_end, 22, 23, font=F_PY, fmt="0.000000", fill=FILL_PY)
    style_range(ws, 5, R_end, 24, 24)
    style_range(ws, 5, R_end, 25, 25, fmt="0.0000")
    style_range(ws, 5, R_end, 26, 26, font=F_PY, fmt="0.000000", fill=FILL_PY)
    ws.cell(R_end + 2, 1, "Excel 與 Python 配對 t p 值之最大絕對差：").font = F_BOLD
    ws.cell(R_end + 2, 5, f"=SUMPRODUCT(MAX(ABS(U5:U{R_end})))").number_format = "0.0E+00"
    ws.cell(R_end + 2, 5).fill = FILL_CHK
    ws.cell(R_end + 3, 1, "說明：主要指標為第 5–20 列；其餘為分節指標。Holm 校正僅針對 16 項主要指標。").font = F_SUB
    ws.freeze_panes = "C5"
    ws.cell(4, 13).comment = Comment("TDIST(|t|, df, 2)，與 Excel 2010 起的 T.DIST.2T 結果相同", "分析程式")

    # ------------------------------------------------------------ 分層檢定
    ws = wb.create_sheet("分層檢定")
    title(ws, "分層配對 t 檢定（主要指標）", "以 SUMIF / SUMPRODUCT 公式依分層條件計算；標準差 = √[(Σd² − n·d̄²)/(n−1)]。")
    HS = ["分層", "指標", "單位", "n", "Σd", "Σd²", "平均差", "差值SD", "SE", "t", "df", "p 值", "95% CI 下限", "95% CI 上限",
          "Python 平均差", "Python 配對 t p", "差異 (p)", "Python 判定 p"]
    head(ws, 4, HS, [12, 24, 8, 6, 10, 10, 10, 10, 9, 8, 5, 10, 11, 11, 11, 11, 10, 11])
    strata_defs = [("sex:M", "男性", SEXR, "M"), ("sex:F", "女性", SEXR, "F"),
                   ("age:<30", "<30 歲", AGER, "<30"), ("age:30-39", "30-39 歲", AGER, "30-39"), ("age:≥40", "≥40 歲", AGER, "≥40")]
    pos = {}  # (stratum, code) -> row
    r = 5
    for key, lab, crit_rng, crit in strata_defs:
        st = res["strata"][key]
        for code in A.MAIN:
            Dd = rng(col[code][2])
            s = st["main"][code] if not st["suppressed"] else {}
            vals = [lab, fields[code][0], fields[code][1], f'=COUNTIF({crit_rng},"={crit}")', f'=SUMIF({crit_rng},"={crit}",{Dd})',
                    f'=SUMPRODUCT(({crit_rng}="{crit}")*({Dd})^2)', f"=E{r}/D{r}", f"=SQRT((F{r}-D{r}*G{r}^2)/(D{r}-1))",
                    f"=H{r}/SQRT(D{r})", f"=IFERROR(G{r}/I{r},0)", f"=D{r}-1", f"=IFERROR(TDIST(ABS(J{r}),K{r},2),1)",
                    f"=G{r}-TINV(0.05,K{r})*I{r}", f"=G{r}+TINV(0.05,K{r})*I{r}",
                    s.get("diff_mean"), s.get("p_t"), f"=L{r}-P{r}" if s else "", s.get("p")]
            for j, v in enumerate(vals, 1):
                ws.cell(r, j, v)
            pos[(key, code)] = r
            r += 1
    style_range(ws, 5, r - 1, 1, 18)
    style_range(ws, 5, r - 1, 5, 14, fmt="0.0000")
    style_range(ws, 5, r - 1, 15, 16, font=F_PY, fmt="0.000000", fill=FILL_PY)
    style_range(ws, 5, r - 1, 17, 17, fmt="0.0E+00", fill=FILL_CHK)
    style_range(ws, 5, r - 1, 18, 18, font=F_PY, fmt="0.000000", fill=FILL_PY)
    ws.freeze_panes = "D5"

    # ------------------------------------------------------------ 組間比較
    ws = wb.create_sheet("組間比較")
    title(ws, "分層之間變化量比較", "性別：Welch t（公式）；年齡層：單因子 ANOVA（公式，引用「分層檢定」的 n、平均、SD）。")
    HC = ["指標", "Welch t", "Welch df", "Welch p", "Python Welch p", "差異", "Python Mann-Whitney p",
          "總平均", "SS 組間", "SS 組內", "F", "df1", "df2", "ANOVA p", "Python ANOVA p", "差異", "Python Kruskal-Wallis p"]
    head(ws, 4, HC, [24, 9, 9, 10, 11, 9, 12, 10, 10, 10, 9, 5, 5, 10, 11, 9, 12])
    SS = "分層檢定"
    r = 5
    for code in A.MAIN:
        m, f = pos[("sex:M", code)], pos[("sex:F", code)]
        g = [pos[(k, code)] for k in ("age:<30", "age:30-39", "age:≥40")]
        cpy = res["compare"][code]
        va = f"({SS}!H{m}^2/{SS}!D{m})"
        vb = f"({SS}!H{f}^2/{SS}!D{f})"
        n_ = [f"{SS}!D{x}" for x in g]
        m_ = [f"{SS}!G{x}" for x in g]
        s_ = [f"{SS}!H{x}" for x in g]
        vals = [fields[code][0],
                f"=({SS}!G{m}-{SS}!G{f})/SQRT({va}+{vb})",
                f"=({va}+{vb})^2/({va}^2/({SS}!D{m}-1)+{vb}^2/({SS}!D{f}-1))",
                f"=BETADIST(C{r}/(C{r}+B{r}^2),C{r}/2,0.5)",
                cpy["sex_p_welch"], f"=D{r}-E{r}", cpy["sex_p_mw"],
                "=(" + "+".join(f"{a}*{b}" for a, b in zip(n_, m_)) + ")/(" + "+".join(n_) + ")",
                "=" + "+".join(f"{a}*({b}-H{r})^2" for a, b in zip(n_, m_)),
                "=" + "+".join(f"({a}-1)*{b}^2" for a, b in zip(n_, s_)),
                f"=(I{r}/L{r})/(J{r}/M{r})", "=3-1", "=" + "+".join(n_) + "-3",
                f"=FDIST(K{r},L{r},M{r})", cpy["age_p_anova"], f"=N{r}-O{r}", cpy["age_p_kw"]]
        for j, v in enumerate(vals, 1):
            ws.cell(r, j, v)
        r += 1
    style_range(ws, 5, r - 1, 1, 17, fmt="0.0000")
    for cc in (5, 7, 15, 17):
        style_range(ws, 5, r - 1, cc, cc, font=F_PY, fmt="0.000000", fill=FILL_PY)
    for cc in (6, 16):
        style_range(ws, 5, r - 1, cc, cc, fmt="0.0E+00", fill=FILL_CHK)
    ws.cell(r + 1, 1, "Welch df 為非整數。Excel 的 TDIST／T.DIST 會把自由度截為整數，故雙尾 p 改以等價式 BETADIST(df/(df+t²), df/2, 0.5) 計算，保留小數自由度（與 SPSS 相同）。").font = F_SUB

    # ------------------------------------------------------------ McNemar
    ws = wb.create_sheet("McNemar")
    title(ws, "風險分類前後變化（完成前後測者）", "a=前後皆是、b=前是後否（改善）、c=前否後是（惡化）、d=前後皆否。精確 p = min(1, 2×BINOM.DIST(min(b,c), b+c, 0.5, 累積))。")
    HM = ["風險分類", "a 前後皆是", "b 改善", "c 惡化", "d 前後皆否", "前測人數", "後測人數", "McNemar 精確 p", "Python p", "差異",
          "McNemar χ²(參考)", "χ² p(參考)"]
    head(ws, 4, HM, [30, 10, 9, 9, 10, 9, 9, 13, 10, 9, 13, 11])
    r = 5
    for bc, (cpre, cpost) in flag_cols.items():
        X0, X1 = rng(cpre), rng(cpost)
        vals = [A.BINARY[bc][0], f"=COUNTIFS({X0},1,{X1},1)", f"=COUNTIFS({X0},1,{X1},0)", f"=COUNTIFS({X0},0,{X1},1)",
                f"=COUNTIFS({X0},0,{X1},0)", f"=B{r}+C{r}", f"=B{r}+D{r}",
                f"=IF(C{r}+D{r}=0,1,MIN(1,2*BINOMDIST(MIN(C{r},D{r}),C{r}+D{r},0.5,TRUE)))",
                res["strata"]["all"]["binary"][bc]["p"], f"=H{r}-I{r}",
                f"=IF(C{r}+D{r}=0,0,(C{r}-D{r})^2/(C{r}+D{r}))", f"=CHIDIST(K{r},1)"]
        for j, v in enumerate(vals, 1):
            ws.cell(r, j, v)
        r += 1
    style_range(ws, 5, r - 1, 1, 12)
    style_range(ws, 5, r - 1, 8, 8, fmt="0.000000")
    style_range(ws, 5, r - 1, 9, 9, font=F_PY, fmt="0.000000", fill=FILL_PY)
    style_range(ws, 5, r - 1, 10, 10, fmt="0.0E+00", fill=FILL_CHK)
    style_range(ws, 5, r - 1, 11, 12, fmt="0.0000")
    # Bowker
    r += 1
    for key, cats, (c0, c1), label in [("bmi", A.BMI_CATS, cat_cols["bmi"], "BMI 分級"),
                                       ("vfat", A.VFAT_CATS, cat_cols["vfat"], "內臟脂肪分級")]:
        ws.cell(r, 1, f"{label}：前測（列）× 後測（欄）交叉表與 Bowker 對稱性檢定").font = F_BOLD
        r += 1
        ws.cell(r, 1, "前測 \\ 後測").font = F_HEAD
        ws.cell(r, 1).fill = FILL_HEAD
        for j, cname in enumerate(cats):
            ws.cell(r, 2 + j, cname).font = F_HEAD
            ws.cell(r, 2 + j).fill = FILL_HEAD
        top = r + 1
        for i, ri in enumerate(cats):
            ws.cell(top + i, 1, ri).font = F_BOLD
            for j, cj in enumerate(cats):
                ws.cell(top + i, 2 + j, f'=COUNTIFS({rng(c0)},"{ri}",{rng(c1)},"{cj}")').font = F_BASE
        k = len(cats)
        terms, dfs = [], []
        for i in range(k):
            for j in range(i + 1, k):
                a_ = f"{L(2 + j)}{top + i}"
                b_ = f"{L(2 + i)}{top + j}"
                terms.append(f"IF({a_}+{b_}>0,({a_}-{b_})^2/({a_}+{b_}),0)")
                dfs.append(f"IF({a_}+{b_}>0,1,0)")
        rr = top + k + 1
        py = res["strata"]["all"]["bmi_cat" if key == "bmi" else "vfat_cat"]
        for j, (lab, v, fmt) in enumerate([("Bowker χ²", "=" + "+".join(terms), "0.0000"), ("自由度", "=" + "+".join(dfs), "0"),
                                           ("p 值", f"=IF(C{rr}=0,1,CHIDIST(B{rr},C{rr}))", "0.000000"),
                                           ("Python p", py["p"], "0.000000"), ("差異", f"=D{rr}-E{rr}", "0.0E+00")]):
            ws.cell(rr - 1, 2 + j, lab).font = F_BOLD
            cell = ws.cell(rr, 2 + j, v)
            cell.number_format = fmt
            cell.font = F_PY if lab == "Python p" else F_BASE
            if lab == "Python p":
                cell.fill = FILL_PY
            if lab == "差異":
                cell.fill = FILL_CHK
        r = rr + 2

    # ------------------------------------------------------------ 卡方檢定
    ws = wb.create_sheet("卡方檢定")
    title(ws, "前測風險分類 × 性別／年齡層：Pearson 卡方檢定", "期望次數 = 列合計 × 欄合計 ÷ 總數；χ² = Σ(O−E)²/E；p = CHIDIST(χ², df)。期望次數 < 5 的格子超過 20% 時，參考 Fisher 精確檢定（Python）。")
    ws.column_dimensions["A"].width = 30
    for j in range(2, 12):
        ws.column_dimensions[L(j)].width = 11
    r = 4
    SEXB, AGEB = brng(2), brng(4)
    for x in res["baseline_chisq"]:
        bc = x["outcome"]
        grp_rng = SEXB if x["by"] == "性別" else AGEB
        keys = ["M", "F"] if x["by"] == "性別" else [g for *_, g in A.AGE_BINS]
        ws.cell(r, 1, f"{x['outcome_name']} × {x['by']}").font = F_BOLD
        r += 1
        for j, v in enumerate(["觀察次數 O", "是", "否", "合計", "", "期望次數 E", "是", "否"], 1):
            ws.cell(r, j, v).font = F_HEAD
            ws.cell(r, j).fill = FILL_HEAD
        top = r + 1
        for i, (lab, kk) in enumerate(zip(x["rows"], keys)):
            rr = top + i
            ws.cell(rr, 1, lab).font = F_BOLD
            ws.cell(rr, 2, f'=COUNTIFS({grp_rng},"={kk}",{brng(fcol[bc])},1)').font = F_BASE
            ws.cell(rr, 3, f'=COUNTIFS({grp_rng},"={kk}",{brng(fcol[bc])},0)').font = F_BASE
            ws.cell(rr, 4, f"=B{rr}+C{rr}").font = F_BASE
            ws.cell(rr, 6, lab).font = F_BOLD
        last = top + len(keys) - 1
        tot = last + 1
        ws.cell(tot, 1, "合計").font = F_BOLD
        for cc in (2, 3, 4):
            ws.cell(tot, cc, f"=SUM({L(cc)}{top}:{L(cc)}{last})").font = F_BOLD
        for i in range(len(keys)):
            rr = top + i
            for cc, oc in ((7, 2), (8, 3)):
                ws.cell(rr, cc, f"=$D{rr}*{L(oc)}${tot}/$D${tot}").number_format = "0.00"
        Orng, Erng = f"B{top}:C{last}", f"G{top}:H{last}"
        sr = tot + 1
        labels = ["χ²", "df", "p 值", "p 驗證 (CHITEST)", "Python p", "差異", "最小期望次數", "期望 < 5 格數", "Python Fisher p"]
        formulas = [f"=SUMPRODUCT(({Orng}-{Erng})^2/{Erng})", f"=({len(keys)}-1)*(2-1)", f"=CHIDIST(B{sr + 1},C{sr + 1})",
                    f"=CHITEST({Orng},{Erng})", x["p"], f"=D{sr + 1}-F{sr + 1}", f"=MIN({Erng})", f'=COUNTIF({Erng},"<5")', x["p_fisher"]]
        for j, (lab, fml) in enumerate(zip(labels, formulas)):
            ws.cell(sr, 2 + j, lab).font = F_SUB
            cell = ws.cell(sr + 1, 2 + j, fml)
            cell.font = F_PY if lab.startswith("Python") else F_BASE
            cell.number_format = "0.000000" if "p" in lab else "0.0E+00" if lab == "差異" else "0.0000"
            if lab.startswith("Python"):
                cell.fill = FILL_PY
            if lab == "差異":
                cell.fill = FILL_CHK
        r = sr + 3

    # ------------------------------------------------------------ 流失分析
    ws = wb.create_sheet("流失分析")
    title(ws, "完成前後測者 vs 僅前測者（前測值）", "Welch t 檢定以 AVERAGEIF / SUMPRODUCT 公式計算；p = BETADIST(df/(df+t²), df/2, 0.5)，保留小數自由度。")
    HA = ["變項", "完成 n", "完成 平均", "完成 SD", "僅前測 n", "僅前測 平均", "僅前測 SD", "Welch t", "Welch df", "p 值", "Python p", "差異"]
    head(ws, 4, HA, [22, 8, 11, 10, 9, 11, 10, 9, 9, 10, 10, 9])
    CMP = brng(5)
    r = 5
    att_cols = {"AGE": 3, **bcol}
    for v in res["attrition"]["vars"]:
        X = brng(att_cols[v["code"]])
        sd = lambda flag, nn, mm: f"=SQRT((SUMPRODUCT(({CMP}={flag})*({X})^2)-{nn}*{mm}^2)/({nn}-1))"
        vals = [v["name"], f"=COUNTIF({CMP},1)", f"=AVERAGEIF({CMP},1,{X})", sd(1, f"B{r}", f"C{r}"),
                f"=COUNTIF({CMP},0)", f"=AVERAGEIF({CMP},0,{X})", sd(0, f"E{r}", f"F{r}"),
                f"=(C{r}-F{r})/SQRT(D{r}^2/B{r}+G{r}^2/E{r})",
                f"=(D{r}^2/B{r}+G{r}^2/E{r})^2/((D{r}^2/B{r})^2/(B{r}-1)+(G{r}^2/E{r})^2/(E{r}-1))",
                f"=BETADIST(I{r}/(I{r}+H{r}^2),I{r}/2,0.5)", v["p"], f"=J{r}-K{r}"]
        for j, val in enumerate(vals, 1):
            ws.cell(r, j, val)
        r += 1
    style_range(ws, 5, r - 1, 1, 12, fmt="0.0000")
    style_range(ws, 5, r - 1, 11, 11, font=F_PY, fmt="0.000000", fill=FILL_PY)
    style_range(ws, 5, r - 1, 12, 12, fmt="0.0E+00", fill=FILL_CHK)
    r += 1
    ws.cell(r, 1, "性別 × 是否完成後測（卡方）").font = F_BOLD
    r += 1
    for j, v in enumerate(["", "完成", "僅前測", "合計", "期望 完成", "期望 僅前測"], 1):
        ws.cell(r, j, v).font = F_HEAD
        ws.cell(r, j).fill = FILL_HEAD
    t0 = r + 1
    for i, (lab, kk) in enumerate([("男性", "M"), ("女性", "F")]):
        rr = t0 + i
        ws.cell(rr, 1, lab).font = F_BOLD
        ws.cell(rr, 2, f'=COUNTIFS({SEXB},"{kk}",{CMP},1)')
        ws.cell(rr, 3, f'=COUNTIFS({SEXB},"{kk}",{CMP},0)')
        ws.cell(rr, 4, f"=B{rr}+C{rr}")
    tt = t0 + 2
    ws.cell(tt, 1, "合計").font = F_BOLD
    for cc in (2, 3, 4):
        ws.cell(tt, cc, f"=SUM({L(cc)}{t0}:{L(cc)}{t0 + 1})")
    for i in range(2):
        rr = t0 + i
        ws.cell(rr, 5, f"=$D{rr}*B${tt}/$D${tt}").number_format = "0.00"
        ws.cell(rr, 6, f"=$D{rr}*C${tt}/$D${tt}").number_format = "0.00"
    sr = tt + 2
    for j, (lab, fml) in enumerate([("χ²", f"=SUMPRODUCT((B{t0}:C{t0 + 1}-E{t0}:F{t0 + 1})^2/E{t0}:F{t0 + 1})"),
                                    ("p 值", f"=CHIDIST(B{sr + 1},1)"), ("Python p", res["attrition"]["sex"]["p"]),
                                    ("差異", f"=C{sr + 1}-D{sr + 1}"), ("Python Fisher p", res["attrition"]["sex"]["p_fisher"])]):
        ws.cell(sr, 2 + j, lab).font = F_SUB
        cell = ws.cell(sr + 1, 2 + j, fml)
        cell.number_format = "0.000000" if lab != "差異" else "0.0E+00"
        cell.font = F_PY if lab.startswith("Python") else F_BASE
        if lab.startswith("Python"):
            cell.fill = FILL_PY
        if lab == "差異":
            cell.fill = FILL_CHK

    for wsx in wb.worksheets:
        wsx.sheet_view.showGridLines = wsx.title not in ("說明",)
    out_path = Path(out_path)
    wb.save(out_path)
    write_sps(out_path, codes, R1, len(hdr2), B1, len(h2))
    return out_path


def write_sps(xlsx, codes, last_row, ncol, last_base_row, nbase):
    main = list(A.MAIN)
    pre_vars = " ".join(f"{c}_pre" for c in main)
    post_vars = " ".join(f"{c}_post" for c in main)
    d_vars = " ".join(f"{c}_d" for c in main)
    flags = list(A.BINARY)
    txt = f"""* SPSS 語法：以 Excel 驗算檔重跑主要檢定，對照報告與 Excel 結果.
* 使用前請將 FILE 路徑改為 Excel 檔實際位置。資料自第 2 列（變項代碼列）讀入.
SET UNICODE=ON.
GET DATA /TYPE=XLSX /FILE='{xlsx.name}' /SHEET=name '配對資料'
  /CELLRANGE=RANGE 'A2:{L(ncol)}{last_row}' /READNAMES=ON /ASSUMEDSTRWIDTH=32767.
DATASET NAME paired WINDOW=FRONT.

* 1. 配對 t 檢定（全體）.
T-TEST PAIRS={pre_vars} WITH {post_vars} (PAIRED)
  /CRITERIA=CI(.9500) /MISSING=ANALYSIS.

* 2. Wilcoxon 符號等級檢定（漸近顯著性）.
NPAR TESTS /WILCOXON={pre_vars} WITH {post_vars} (PAIRED) /MISSING ANALYSIS.

* 3. 差值常態性（Shapiro-Wilk）.
EXAMINE VARIABLES={d_vars} /PLOT NPPLOT /STATISTICS NONE.

* 4. 分層配對 t 檢定.
SORT CASES BY SEX.
SPLIT FILE LAYERED BY SEX.
T-TEST PAIRS={pre_vars} WITH {post_vars} (PAIRED).
SPLIT FILE OFF.
SORT CASES BY AGEGRP.
SPLIT FILE LAYERED BY AGEGRP.
T-TEST PAIRS={pre_vars} WITH {post_vars} (PAIRED).
SPLIT FILE OFF.

* 5. 組間比較變化量：性別（Welch 見「不假設變異數相等」列）、年齡層.
T-TEST GROUPS=SEX('M' 'F') /VARIABLES={d_vars}.
NPAR TESTS /M-W={d_vars} BY SEX('M' 'F').
AUTORECODE VARIABLES=AGEGRP /INTO AGEGRP_N.
ONEWAY {d_vars} BY AGEGRP_N.
NPAR TESTS /K-W={d_vars} BY AGEGRP_N(1 3).

* 6. McNemar（風險分類前後）.
CROSSTABS /TABLES={' '.join(f'{f}_pre' for f in flags)} BY {' '.join(f'{f}_post' for f in flags)}
  /STATISTICS=MCNEMAR /CELLS=COUNT.
CROSSTABS /TABLES=BMIcat_pre BY BMIcat_post /STATISTICS=MCNEMAR /CELLS=COUNT.

* 7. 前測卡方檢定（全部前測受測者）.
GET DATA /TYPE=XLSX /FILE='{xlsx.name}' /SHEET=name '前測資料'
  /CELLRANGE=RANGE 'A2:{L(nbase)}{last_base_row}' /READNAMES=ON /ASSUMEDSTRWIDTH=32767.
DATASET NAME baseline WINDOW=FRONT.
CROSSTABS /TABLES={' '.join(flags)} BY SEX AGEGRP
  /STATISTICS=CHISQ /CELLS=COUNT EXPECTED ROW /METHOD=EXACT TIMER(5).
CROSSTABS /TABLES=SEX BY COMPLETED /STATISTICS=CHISQ /CELLS=COUNT EXPECTED.
"""
    xlsx.with_suffix(".sps").write_text(txt, encoding="utf-8")
